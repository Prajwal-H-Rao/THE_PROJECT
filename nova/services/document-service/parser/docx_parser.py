"""DOCX parser using python-docx with paragraph provenance."""

from pathlib import Path

from docx import Document

from nova.services.document_service.parser.base import BaseParser, ParseResult
from nova.services.document_service.schemas import ParserSettings
from nova.services.document_service.schemas.extraction import (
    Confidence,
    FieldExtraction,
    ParsedDocument,
    ParsedDocumentMetadata,
    ReviewStatus,
    SourceSpan,
)


class DOCXParser(BaseParser):
    """DOCX parser using python-docx with paragraph provenance."""

    def __init__(self, settings: ParserSettings):
        super().__init__(settings)
        self.parser_version = "python-docx"

    def parse(self, file_path: Path, mime_type: str) -> ParseResult:
        """Parse a DOCX file and extract content with paragraph provenance."""
        self.warnings = []
        doc = Document(file_path)

        try:
            fields = []
            raw_text_parts = []
            structure = {"paragraphs": [], "tables": []}
            char_offset = 0

            # Extract paragraphs
            for para_idx, para in enumerate(doc.paragraphs):
                text = para.text
                if not text.strip():
                    continue

                raw_text_parts.append(text)
                source_ref = SourceSpan(
                    line=para_idx + 1,
                    char_start=char_offset,
                    char_end=char_offset + len(text),
                    raw_text=text[:500],
                )

                field = self._create_field(
                    field_name=f"paragraph_{para_idx}",
                    value=text,
                    source_ref=source_ref,
                    confidence=Confidence(value=1.0, method="parser"),
                    review_status=ReviewStatus.REVIEWED,
                    metadata={
                        "style": para.style.name if para.style else "Normal",
                        "alignment": str(para.alignment) if para.alignment else None,
                    },
                )
                fields.append(field)

                structure["paragraphs"].append({
                    "index": para_idx,
                    "style": para.style.name if para.style else "Normal",
                    "text_preview": text[:200],
                })
                char_offset += len(text) + 1  # +1 for newline

            # Extract tables
            for table_idx, table in enumerate(doc.tables):
                table_data = []
                for row_idx, row in enumerate(table.rows):
                    row_data = []
                    for cell_idx, cell in enumerate(row.cells):
                        cell_text = cell.text
                        if cell_text.strip():
                            source_ref = SourceSpan(
                                line=row_idx + 1,
                                column=cell_idx + 1,
                                raw_text=cell_text[:200],
                            )
                            field = self._create_field(
                                field_name=f"table_{table_idx}_row_{row_idx}_cell_{cell_idx}",
                                value=cell_text,
                                source_ref=source_ref,
                                confidence=Confidence(value=1.0, method="parser"),
                                review_status=ReviewStatus.REVIEWED,
                                metadata={"table": table_idx, "row": row_idx, "cell": cell_idx},
                            )
                            fields.append(field)
                        row_data.append(cell_text)
                    table_data.append(row_data)

                structure["tables"].append({
                    "index": table_idx,
                    "rows": len(table.rows),
                    "cols": len(table.columns),
                    "data": table_data,
                })

            # Core properties
            core_props = doc.core_properties
            for attr in ["title", "author", "created", "modified", "subject", "keywords", "category"]:
                value = getattr(core_props, attr, None)
                if value:
                    fields.append(self._create_field(
                        field_name=f"docx_meta_{attr}",
                        value=str(value),
                        confidence=Confidence(value=1.0, method="parser"),
                        metadata={"source": "docx_core_properties"},
                    ))

            metadata = ParsedDocumentMetadata(
                document_id="",
                original_filename=file_path.name,
                mime_type=mime_type,
                file_size=file_path.stat().st_size,
                sha256="",
                page_count=None,
                parser_version=self.parser_version,
                processing_time_ms=0,
            )

            document = ParsedDocument(
                metadata=metadata,
                fields=fields,
                raw_content="\n".join(raw_text_parts),
                structure=structure,
                warnings=self.warnings.copy(),
            )

            return ParseResult(document=document, warnings=self.warnings.copy())

        except Exception as e:
            self._add_warning(f"DOCX parsing error: {e}")
            raise